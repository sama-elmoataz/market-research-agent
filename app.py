import io
import re
import uuid
from html import escape
from pathlib import Path

import streamlit as st

from PIL import Image

from docx import Document

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from langgraph.types import Command

from src.graph import research_graph


# =========================================================
# PATHS
# =========================================================

BASE_DIR = Path(__file__).resolve().parent

LOGO_PATH = (
    BASE_DIR
    / "assets"
    / "logo.png"
)


# =========================================================
# LOGO PROCESSING
# =========================================================

def load_clean_logo(
    image_path: Path,
):
    """
    Remove the dark background from the logo
    and crop unnecessary empty space.
    """

    image = Image.open(
        image_path
    ).convert("RGBA")

    pixels = image.load()

    width, height = image.size

    # Use the top-left pixel as the
    # approximate background color.
    bg_r, bg_g, bg_b, _ = pixels[0, 0]

    for y in range(height):
        for x in range(width):

            r, g, b, a = pixels[x, y]

            distance = (
                (r - bg_r) ** 2
                + (g - bg_g) ** 2
                + (b - bg_b) ** 2
            ) ** 0.5

            # Remove pixels close to
            # the original dark background.
            if distance < 45:

                pixels[x, y] = (
                    r,
                    g,
                    b,
                    0,
                )

    # Crop transparent empty space.
    bbox = image.getbbox()

    if bbox:

        image = image.crop(
            bbox
        )

    return image


CLEAN_LOGO = (
    load_clean_logo(
        LOGO_PATH
    )
    if LOGO_PATH.exists()
    else None
)


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="Market Research Agent",
    page_icon=str(LOGO_PATH),
    layout="wide",
    initial_sidebar_state="collapsed",
)


# =========================================================
# SESSION STATE
# =========================================================

DEFAULTS = {
    "thread_id": None,
    "topic": "",
    "workflow_status": "idle",
    "graph_state": {},
    "report": "",
    "error_message": "",
    "mode": "New Research",
}


for key, value in DEFAULTS.items():

    if key not in st.session_state:
        st.session_state[key] = value


# =========================================================
# LANGGRAPH CONFIGURATION
# =========================================================

def build_config(
    thread_id: str,
    run_mode: str,
):
    return {
        "configurable": {
            "thread_id": thread_id,
        },

        "tags": [
            "market-research",
            "streamlit",
        ],

        "metadata": {
            "thread_id": thread_id,
            "run_mode": run_mode,
            "interface": "streamlit",
        },
    }


def build_initial_state(
    topic: str,
):
    return {
        "topic": topic,

        "research_plan": [],

        "research_results": [],

        "research_quality_status": "",

        "research_quality_reasoning": "",

        "research_gaps": [],

        "additional_queries": [],

        "research_round": 0,

        "analysis": "",

        "draft_report": "",

        "critique": "",

        "review_status": "",

        "iteration": 0,

        "approved": False,
    }


# =========================================================
# SNAPSHOT / STATE SYNC
# =========================================================

def sync_snapshot(
    thread_id: str,
):
    config = build_config(
        thread_id=thread_id,
        run_mode="resume",
    )

    snapshot = research_graph.get_state(
        config
    )

    if not snapshot.values:

        raise ValueError(
            "No checkpoint was found "
            "for this thread ID."
        )

    state = dict(
        snapshot.values
    )

    st.session_state.graph_state = state

    st.session_state.topic = state.get(
        "topic",
        st.session_state.topic,
    )

    st.session_state.report = state.get(
        "draft_report",
        "",
    )

    # -----------------------------------------
    # Determine workflow status
    # -----------------------------------------

    if not snapshot.next:

        st.session_state.workflow_status = (
            "completed"
        )

    elif (
        "human_approval"
        in snapshot.next
    ):

        st.session_state.workflow_status = (
            "awaiting_approval"
        )

    else:

        st.session_state.workflow_status = (
            "paused"
        )

    return snapshot


# =========================================================
# RESULT PROCESSING
# =========================================================

def process_graph_result(
    result,
    thread_id: str,
):
    sync_snapshot(
        thread_id
    )

    if (
        isinstance(result, dict)
        and "__interrupt__" in result
    ):

        interrupts = result.get(
            "__interrupt__",
            [],
        )

        if interrupts:

            interrupt_value = (
                interrupts[0].value
            )

            if isinstance(
                interrupt_value,
                dict,
            ):

                interrupt_report = (
                    interrupt_value.get(
                        "report",
                        "",
                    )
                )

                if interrupt_report:

                    st.session_state.report = (
                        interrupt_report
                    )

        st.session_state.workflow_status = (
            "awaiting_approval"
        )


# =========================================================
# START NEW RESEARCH
# =========================================================

def start_research(
    topic: str,
):
    thread_id = str(
        uuid.uuid4()
    )

    st.session_state.thread_id = (
        thread_id
    )

    st.session_state.topic = (
        topic
    )

    st.session_state.workflow_status = (
        "running"
    )

    st.session_state.error_message = ""

    config = build_config(
        thread_id=thread_id,
        run_mode="new",
    )

    initial_state = build_initial_state(
        topic
    )

    result = research_graph.invoke(
        initial_state,
        config=config,
    )

    process_graph_result(
        result=result,
        thread_id=thread_id,
    )


# =========================================================
# LOAD EXISTING THREAD
# =========================================================

def load_existing_thread(
    thread_id: str,
):
    if not thread_id:

        raise ValueError(
            "Enter a thread ID."
        )

    st.session_state.thread_id = (
        thread_id
    )

    st.session_state.error_message = ""

    sync_snapshot(
        thread_id
    )


# =========================================================
# RESUME NORMAL WORKFLOW
# =========================================================

def resume_workflow():

    thread_id = (
        st.session_state.thread_id
    )

    if not thread_id:

        raise ValueError(
            "No active research thread."
        )

    st.session_state.workflow_status = (
        "running"
    )

    config = build_config(
        thread_id=thread_id,
        run_mode="resume",
    )

    result = research_graph.invoke(
        None,
        config=config,
    )

    process_graph_result(
        result=result,
        thread_id=thread_id,
    )


# =========================================================
# HUMAN APPROVAL
# =========================================================

def submit_human_decision(
    decision: str,
):

    thread_id = (
        st.session_state.thread_id
    )

    if not thread_id:

        raise ValueError(
            "No active research thread."
        )

    config = build_config(
        thread_id=thread_id,
        run_mode="resume",
    )

    result = research_graph.invoke(
        Command(
            resume=decision
        ),
        config=config,
    )

    process_graph_result(
        result=result,
        thread_id=thread_id,
    )


# =========================================================
# RESET
# =========================================================

def reset_session():

    st.session_state.clear()

    st.rerun()


# =========================================================
# EXPORT HELPERS
# =========================================================

def clean_markdown_text(
    markdown_text: str,
) -> str:
    """
    Convert Markdown report into a cleaner
    plain-text representation.
    """

    text = markdown_text

    # Remove bold / italic markers.
    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text,
    )

    text = re.sub(
        r"\*(.*?)\*",
        r"\1",
        text,
    )

    # Markdown heading markers.
    text = re.sub(
        r"^#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE,
    )

    # Markdown links.
    text = re.sub(
        r"\[([^\]]+)\]\([^)]+\)",
        r"\1",
        text,
    )

    return text.strip()


def build_docx(
    report: str,
) -> bytes:
    """
    Convert the Markdown report into
    a simple professional DOCX document.
    """

    document = Document()

    document.add_heading(
        "Market Research Report",
        level=0,
    )

    if st.session_state.topic:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            st.session_state.topic
        )

        run.italic = True

    for raw_line in report.splitlines():

        line = raw_line.strip()

        if not line:

            continue

        # -------------------------------------
        # Markdown headings
        # -------------------------------------

        if line.startswith("### "):

            document.add_heading(
                line[4:],
                level=3,
            )

        elif line.startswith("## "):

            document.add_heading(
                line[3:],
                level=2,
            )

        elif line.startswith("# "):

            document.add_heading(
                line[2:],
                level=1,
            )

        # -------------------------------------
        # Bullet points
        # -------------------------------------

        elif line.startswith("- "):

            clean_line = clean_markdown_text(
                line[2:]
            )

            document.add_paragraph(
                clean_line,
                style="List Bullet",
            )

        # -------------------------------------
        # Numbered items
        # -------------------------------------

        elif re.match(
            r"^\d+\.\s",
            line,
        ):

            clean_line = re.sub(
                r"^\d+\.\s*",
                "",
                line,
            )

            clean_line = clean_markdown_text(
                clean_line
            )

            document.add_paragraph(
                clean_line,
                style="List Number",
            )

        # -------------------------------------
        # Ignore markdown separators
        # -------------------------------------

        elif line in {
            "---",
            "***",
            "___",
        }:

            continue

        # -------------------------------------
        # Normal paragraph
        # -------------------------------------

        else:

            document.add_paragraph(
                clean_markdown_text(
                    line
                )
            )

    buffer = io.BytesIO()

    document.save(
        buffer
    )

    buffer.seek(0)

    return buffer.getvalue()


def build_pdf(
    report: str,
) -> bytes:
    """
    Convert report into a PDF file.

    Formatting is intentionally clean and
    simple so the export stays reliable.
    """

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Market Research Report",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=20,
        leading=25,
        spaceAfter=14,
    )

    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        spaceAfter=18,
        alignment=TA_CENTER,
    )

    h1_style = ParagraphStyle(
        "CustomHeading1",
        parent=styles["Heading1"],
        fontSize=16,
        leading=20,
        spaceBefore=12,
        spaceAfter=8,
    )

    h2_style = ParagraphStyle(
        "CustomHeading2",
        parent=styles["Heading2"],
        fontSize=13,
        leading=17,
        spaceBefore=10,
        spaceAfter=6,
    )

    h3_style = ParagraphStyle(
        "CustomHeading3",
        parent=styles["Heading3"],
        fontSize=11,
        leading=15,
        spaceBefore=8,
        spaceAfter=5,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontSize=9.5,
        leading=14,
        spaceAfter=7,
    )

    bullet_style = ParagraphStyle(
        "ReportBullet",
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-7,
    )

    story = []

    story.append(
        Paragraph(
            "Market Research Report",
            title_style,
        )
    )

    if st.session_state.topic:

        story.append(
            Paragraph(
                escape(
                    st.session_state.topic
                ),
                subtitle_style,
            )
        )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    for raw_line in report.splitlines():

        line = raw_line.strip()

        if not line:

            continue

        # -------------------------------------
        # Heading 3
        # -------------------------------------

        if line.startswith("### "):

            story.append(
                Paragraph(
                    escape(
                        clean_markdown_text(
                            line[4:]
                        )
                    ),
                    h3_style,
                )
            )

        # -------------------------------------
        # Heading 2
        # -------------------------------------

        elif line.startswith("## "):

            story.append(
                Paragraph(
                    escape(
                        clean_markdown_text(
                            line[3:]
                        )
                    ),
                    h2_style,
                )
            )

        # -------------------------------------
        # Heading 1
        # -------------------------------------

        elif line.startswith("# "):

            story.append(
                Paragraph(
                    escape(
                        clean_markdown_text(
                            line[2:]
                        )
                    ),
                    h1_style,
                )
            )

        # -------------------------------------
        # Bullets
        # -------------------------------------

        elif line.startswith("- "):

            text = clean_markdown_text(
                line[2:]
            )

            story.append(
                Paragraph(
                    "• "
                    + escape(text),
                    bullet_style,
                )
            )

        # -------------------------------------
        # Separators
        # -------------------------------------

        elif line in {
            "---",
            "***",
            "___",
        }:

            story.append(
                Spacer(
                    1,
                    3 * mm,
                )
            )

        # -------------------------------------
        # Markdown table row
        # -------------------------------------

        elif (
            line.startswith("|")
            and line.endswith("|")
        ):

            if re.match(
                r"^\|[\s\-:|]+\|$",
                line,
            ):

                continue

            cells = [
                cell.strip()
                for cell in line.strip(
                    "|"
                ).split("|")
            ]

            table_text = (
                "  |  ".join(
                    cells
                )
            )

            story.append(
                Paragraph(
                    escape(
                        clean_markdown_text(
                            table_text
                        )
                    ),
                    body_style,
                )
            )

        # -------------------------------------
        # Standard paragraph
        # -------------------------------------

        else:

            story.append(
                Paragraph(
                    escape(
                        clean_markdown_text(
                            line
                        )
                    ),
                    body_style,
                )
            )

    document.build(
        story
    )

    buffer.seek(0)

    return buffer.getvalue()


def build_filename(
    extension: str,
) -> str:
    """
    Produce a clean file name based on
    the research topic.
    """

    topic = (
        st.session_state.topic
        or "market_research"
    )

    slug = re.sub(
        r"[^a-zA-Z0-9]+",
        "_",
        topic.lower(),
    )

    slug = slug.strip("_")

    if len(slug) > 55:

        slug = slug[:55].rstrip("_")

    if not slug:

        slug = "market_research"

    return (
        f"{slug}_report.{extension}"
    )


# =========================================================
# HEADER
# =========================================================

header_logo, header_title = (
    st.columns(
        [0.9, 9.1],
        vertical_alignment="center",
    )
)


with header_logo:

    if CLEAN_LOGO is not None:

        st.image(
            CLEAN_LOGO,
            width=100,
        )


with header_title:

    st.title(
        "Market Research Agent"
    )

    st.caption(
        "Autonomous market intelligence "
        "from evidence collection to "
        "human-reviewed insight."
    )


st.write("")


# =========================================================
# CAPABILITY STRIP
# =========================================================

capabilities = [
    (
        "Research Planning",
        "Break down complex questions",
    ),
    (
        "Web Evidence",
        "Collect current public data",
    ),
    (
        "Quality Control",
        "Verify coverage and gaps",
    ),
    (
        "Market Analysis",
        "Turn evidence into insight",
    ),
    (
        "Report Review",
        "Critique and refine content",
    ),
    (
        "Human Approval",
        "Final decision before delivery",
    ),
]


capability_columns = st.columns(
    len(capabilities)
)


for column, (
    capability,
    description,
) in zip(
    capability_columns,
    capabilities,
):

    with column:

        st.markdown(
            f"**{capability}**"
        )

        st.caption(
            description
        )


st.divider()


# =========================================================
# MODE SELECTOR
# =========================================================

if st.session_state.thread_id:

    mode_col, reset_col = st.columns(
        [5, 1]
    )

else:

    mode_col = st.container()

    reset_col = None


with mode_col:

    st.radio(
        "Workflow mode",
        options=[
            "New Research",
            "Resume Research",
        ],
        horizontal=True,
        label_visibility="collapsed",
        key="mode",
    )


if reset_col is not None:

    with reset_col:

        if st.button(
            "New Session",
            width="stretch",
        ):

            reset_session()


# =========================================================
# NEW RESEARCH VIEW
# =========================================================

if (
    st.session_state.mode
    == "New Research"
):

    st.write("")

    outer_left, main_column, outer_right = (
        st.columns(
            [0.45, 8.1, 0.45]
        )
    )


    with main_column:

        with st.container(
            border=True
        ):

            st.subheader(
                "Start a market research project"
            )

            st.caption(
                "Describe the market, geography, "
                "business context, and the questions "
                "you want the agent to investigate."
            )

            st.write("")


            with st.form(
                "research_form"
            ):

                objective = st.text_area(
                    "Research objective",
                    height=220,
                    placeholder=(
                        "Example: Research the premium "
                        "fragrance market in Egypt for "
                        "a new local brand, focusing on "
                        "market growth, competitors, "
                        "consumer preferences, pricing, "
                        "distribution, and opportunities."
                    ),
                    label_visibility="collapsed",
                )


                button_space, button_column = (
                    st.columns(
                        [3.5, 1]
                    )
                )


                with button_column:

                    start_clicked = (
                        st.form_submit_button(
                            "Start Research →",
                            type="primary",
                            width="stretch",
                        )
                    )


            if start_clicked:

                clean_objective = (
                    objective.strip()
                )


                if not clean_objective:

                    st.warning(
                        "Enter a research objective first."
                    )


                else:

                    try:

                        with st.status(
                            "Research workflow in progress",
                            expanded=True,
                        ) as workflow_status:

                            st.write(
                                "Planning the research scope..."
                            )

                            st.write(
                                "Collecting market evidence..."
                            )

                            st.write(
                                "Evaluating research quality..."
                            )

                            st.write(
                                "Analyzing market findings..."
                            )

                            st.write(
                                "Writing and reviewing the report..."
                            )


                            start_research(
                                clean_objective
                            )


                            if (
                                st.session_state.workflow_status
                                == "awaiting_approval"
                            ):

                                workflow_status.update(
                                    label=(
                                        "Research complete — "
                                        "human review required"
                                    ),
                                    state="complete",
                                    expanded=False,
                                )


                            else:

                                workflow_status.update(
                                    label=(
                                        "Research workflow complete"
                                    ),
                                    state="complete",
                                    expanded=False,
                                )


                    except Exception as exc:

                        st.session_state.workflow_status = (
                            "error"
                        )

                        st.session_state.error_message = (
                            f"{type(exc).__name__}: {exc}"
                        )

                        st.error(
                            st.session_state.error_message
                        )


# =========================================================
# RESUME VIEW
# =========================================================

else:

    st.write("")

    outer_left, main_column, outer_right = (
        st.columns(
            [1.2, 6.2, 1.2]
        )
    )


    with main_column:

        with st.container(
            border=True
        ):

            st.subheader(
                "Resume an existing research run"
            )

            st.caption(
                "Load a persistent LangGraph checkpoint "
                "using the thread ID generated when the "
                "research session started."
            )

            st.write("")


            with st.form(
                "resume_form"
            ):

                thread_id_input = (
                    st.text_input(
                        "Thread ID",
                        placeholder=(
                            "Example: "
                            "55612c8b-b1e7-4273-..."
                        ),
                    )
                )


                load_clicked = (
                    st.form_submit_button(
                        "Load Research",
                        type="primary",
                        width="stretch",
                    )
                )


            if load_clicked:

                clean_thread_id = (
                    thread_id_input.strip()
                )


                if not clean_thread_id:

                    st.warning(
                        "Enter a thread ID first."
                    )


                else:

                    try:

                        load_existing_thread(
                            clean_thread_id
                        )

                        st.success(
                            "Research checkpoint loaded."
                        )

                        st.rerun()


                    except Exception as exc:

                        st.error(
                            f"{type(exc).__name__}: {exc}"
                        )


# =========================================================
# CURRENT STATE
# =========================================================

status = (
    st.session_state.workflow_status
)

state = (
    st.session_state.graph_state
)


# =========================================================
# ERROR
# =========================================================

if status == "error":

    st.write("")

    st.error(
        st.session_state.error_message
        or "The research workflow encountered an error."
    )


# =========================================================
# WORKFLOW PREVIEW
# =========================================================

if status == "idle":

    st.write("")
    st.write("")

    st.subheader(
        "How the research workflow works"
    )

    st.caption(
        "A bounded multi-stage agent workflow "
        "combining autonomous research "
        "with deterministic controls."
    )

    st.write("")


    workflow_steps = [
        (
            "01",
            "Plan",
            "Break the objective into "
            "focused research tasks.",
        ),
        (
            "02",
            "Research",
            "Collect current evidence "
            "from public web sources.",
        ),
        (
            "03",
            "Verify",
            "Assess evidence quality "
            "and identify research gaps.",
        ),
        (
            "04",
            "Analyze",
            "Convert research evidence "
            "into market insights.",
        ),
        (
            "05",
            "Review",
            "Critique and revise the report "
            "when necessary.",
        ),
        (
            "06",
            "Approve",
            "Pause for the final "
            "human decision.",
        ),
    ]


    workflow_columns = st.columns(
        6
    )


    for column, (
        number,
        title,
        description,
    ) in zip(
        workflow_columns,
        workflow_steps,
    ):

        with column:

            st.caption(
                number
            )

            st.markdown(
                f"**{title}**"
            )

            st.caption(
                description
            )


# =========================================================
# RESEARCH RUN OVERVIEW
# =========================================================

if (
    status not in [
        "idle",
        "error",
    ]
    and state
):

    st.write("")
    st.write("")

    heading_col, thread_col = (
        st.columns(
            [4.3, 1.7]
        )
    )


    with heading_col:

        st.subheader(
            "Research Run"
        )

        st.caption(
            state.get(
                "topic",
                "",
            )
        )


    with thread_col:

        st.caption(
            "THREAD ID"
        )

        st.code(
            st.session_state.thread_id,
            language=None,
        )


    quality_status = (
        state.get(
            "research_quality_status",
            "",
        )
        or "Pending"
    )

    research_rounds = state.get(
        "research_round",
        0,
    )

    revisions = state.get(
        "iteration",
        0,
    )

    approved = state.get(
        "approved",
        False,
    )


    metric_1, metric_2, metric_3, metric_4 = (
        st.columns(4)
    )


    with metric_1:

        with st.container(
            border=True
        ):

            st.caption(
                "RESEARCH QUALITY"
            )

            st.subheader(
                quality_status
            )


    with metric_2:

        with st.container(
            border=True
        ):

            st.caption(
                "RESEARCH ROUNDS"
            )

            st.subheader(
                str(
                    research_rounds
                )
            )


    with metric_3:

        with st.container(
            border=True
        ):

            st.caption(
                "REPORT REVISIONS"
            )

            st.subheader(
                str(
                    revisions
                )
            )


    with metric_4:

        with st.container(
            border=True
        ):

            st.caption(
                "HUMAN APPROVAL"
            )


            if (
                status
                == "awaiting_approval"
            ):

                approval_text = (
                    "Pending"
                )


            elif approved:

                approval_text = (
                    "Approved"
                )


            elif status == "completed":

                approval_text = (
                    "Not approved"
                )


            else:

                approval_text = (
                    "Pending"
                )


            st.subheader(
                approval_text
            )


# =========================================================
# PAUSED CHECKPOINT
# =========================================================

if status == "paused":

    st.write("")


    with st.container(
        border=True
    ):

        title_col, button_col = (
            st.columns(
                [4, 1]
            )
        )


        with title_col:

            st.subheader(
                "Unfinished research run"
            )

            st.caption(
                "A persistent checkpoint was found. "
                "Continue execution from the saved state."
            )


        with button_col:

            if st.button(
                "Continue Research",
                type="primary",
                width="stretch",
            ):

                try:

                    with st.status(
                        "Resuming research workflow...",
                        expanded=True,
                    ) as resume_status:

                        resume_workflow()

                        resume_status.update(
                            label=(
                                "Workflow resumed successfully"
                            ),
                            state="complete",
                            expanded=False,
                        )

                    st.rerun()


                except Exception as exc:

                    st.error(
                        f"{type(exc).__name__}: {exc}"
                    )


# =========================================================
# REPORT
# =========================================================

report = (
    st.session_state.report
)


if report:

    st.write("")
    st.write("")

    st.subheader(
        "Research Report"
    )

    st.caption(
        st.session_state.topic
    )


    with st.container(
        border=True
    ):

        st.markdown(
            report
        )


# =========================================================
# HUMAN APPROVAL
# =========================================================

if (
    status == "awaiting_approval"
    and report
):

    st.write("")


    with st.container(
        border=True
    ):

        st.subheader(
            "Final Review"
        )

        st.caption(
            "The autonomous workflow has completed "
            "its research and review cycle. "
            "Review the report before making "
            "the final human decision."
        )

        st.write("")


        approve_col, reject_col, empty_col = (
            st.columns(
                [1, 1, 3.5]
            )
        )


        with approve_col:

            if st.button(
                "Approve Report",
                type="primary",
                width="stretch",
            ):

                try:

                    submit_human_decision(
                        "approve"
                    )

                    st.rerun()


                except Exception as exc:

                    st.error(
                        f"{type(exc).__name__}: {exc}"
                    )


        with reject_col:

            if st.button(
                "Reject",
                width="stretch",
            ):

                try:

                    submit_human_decision(
                        "reject"
                    )

                    st.rerun()


                except Exception as exc:

                    st.error(
                        f"{type(exc).__name__}: {exc}"
                    )


# =========================================================
# APPROVED REPORT DOWNLOAD CENTER
# =========================================================

if (
    status == "completed"
    and state
    and state.get(
        "approved",
        False,
    )
    and report
):

    st.write("")
    st.write("")


    with st.container(
        border=True
    ):

        st.subheader(
            "Approved Report"
        )

        st.caption(
            "The research workflow is complete. "
            "Export the approved report in "
            "your preferred format."
        )

        st.write("")


        try:

            txt_report = (
                clean_markdown_text(
                    report
                )
            )

            docx_report = (
                build_docx(
                    report
                )
            )

            pdf_report = (
                build_pdf(
                    report
                )
            )


            md_col, pdf_col, docx_col, txt_col = (
                st.columns(4)
            )


            with md_col:

                st.download_button(
                    label="Markdown",
                    data=report,
                    file_name=build_filename(
                        "md"
                    ),
                    mime="text/markdown",
                    width="stretch",
                    key="download_md",
                )


            with pdf_col:

                st.download_button(
                    label="PDF",
                    data=pdf_report,
                    file_name=build_filename(
                        "pdf"
                    ),
                    mime="application/pdf",
                    width="stretch",
                    key="download_pdf",
                )


            with docx_col:

                st.download_button(
                    label="Word",
                    data=docx_report,
                    file_name=build_filename(
                        "docx"
                    ),
                    mime=(
                        "application/"
                        "vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                    width="stretch",
                    key="download_docx",
                )


            with txt_col:

                st.download_button(
                    label="Plain Text",
                    data=txt_report,
                    file_name=build_filename(
                        "txt"
                    ),
                    mime="text/plain",
                    width="stretch",
                    key="download_txt",
                )


        except Exception as exc:

            st.error(
                "The report was approved, "
                "but one of the export formats "
                f"could not be generated: {exc}"
            )


# =========================================================
# COMPLETED BUT REJECTED
# =========================================================

if (
    status == "completed"
    and state
    and not state.get(
        "approved",
        False,
    )
):

    st.write("")

    st.warning(
        "Research completed without human approval."
    )


# =========================================================
# QUALITY DETAILS
# =========================================================

if state:

    st.write("")


    with st.expander(
        "Research quality details"
    ):

        reasoning = state.get(
            "research_quality_reasoning",
            "",
        )

        research_gaps = state.get(
            "research_gaps",
            [],
        )


        if reasoning:

            st.markdown(
                "**Quality assessment**"
            )

            st.write(
                reasoning
            )


        if research_gaps:

            st.write("")

            st.markdown(
                "**Remaining evidence gaps**"
            )

            for gap in research_gaps:

                st.write(
                    f"• {gap}"
                )


        if (
            not reasoning
            and not research_gaps
        ):

            st.caption(
                "No additional quality details."
            )


# =========================================================
# RUN INFORMATION
# =========================================================

if state:

    with st.expander(
        "Run information"
    ):

        info_left, info_right = (
            st.columns(2)
        )


        with info_left:

            st.caption(
                "THREAD ID"
            )

            st.code(
                st.session_state.thread_id,
                language=None,
            )


            st.caption(
                "REVIEW STATUS"
            )

            st.write(
                state.get(
                    "review_status",
                    "Not available",
                )
            )


        with info_right:

            st.caption(
                "SUPPLEMENTAL RESEARCH ROUNDS"
            )

            st.write(
                state.get(
                    "research_round",
                    0,
                )
            )


            st.caption(
                "REPORT REVISIONS"
            )

            st.write(
                state.get(
                    "iteration",
                    0,
                )
            )